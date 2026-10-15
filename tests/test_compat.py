import io
import os
from pathlib import Path

from openpyxl import load_workbook

from centreon_inventory import compat
from centreon_inventory.excel import write_excel
from centreon_inventory.matrix import build_matrix
from centreon_inventory.sources.clapi_export import parse_clapi_export

SAMPLE = Path(__file__).resolve().parent.parent / "examples" / "centreon_export_exemple.txt"


def test_replace_file_retries_when_target_locked(tmp_path, monkeypatch):
    src, dst = tmp_path / "a.tmp", tmp_path / "a.json"
    src.write_text("{}", encoding="utf-8")
    real, calls = os.replace, []

    def flaky(a, b):  # simule le verrou Windows sur les deux premiers essais
        calls.append(1)
        if len(calls) < 3:
            raise PermissionError("verrouillé")
        real(a, b)
    monkeypatch.setattr(compat.os, "replace", flaky)
    compat.replace_file(src, dst, delay=0)
    assert dst.read_text(encoding="utf-8") == "{}" and len(calls) == 3


def test_write_excel_to_stream():
    buf = io.BytesIO()
    write_excel(build_matrix(parse_clapi_export(SAMPLE)), buf)
    buf.seek(0)
    assert "Synthèse" in load_workbook(buf).sheetnames


def test_cli_output_with_legacy_console_encoding(tmp_path):
    import subprocess
    import sys
    env = dict(os.environ, PYTHONIOENCODING="cp1252", PYTHONUTF8="0")
    r = subprocess.run([sys.executable, "-m", "centreon_inventory", "clapi", str(SAMPLE), "--client", "Réseau → test",
                        "--client-hotes", "sw-", "-o", str(tmp_path / "m.xlsx")],
                       capture_output=True, env=env, cwd=Path(__file__).resolve().parent.parent)
    assert r.returncode == 0, r.stderr.decode("cp1252", "replace")
    assert (tmp_path / "m.xlsx").exists()
