from pathlib import Path

from openpyxl import load_workbook

from centreon_inventory.mapping import Mapper
from centreon_inventory.matrix import HERITE, MIGRER_NATIF, NE_PAS_MIGRER, build_matrix, mask_secrets
from centreon_inventory.sources.clapi_export import parse_clapi_export, parse_clapi_lines
from centreon_inventory.sources.database import build_inventory
from centreon_inventory.excel import write_csv, write_excel

SAMPLE = Path(__file__).resolve().parent.parent / "examples" / "centreon_export_exemple.txt"


def _rows(rows, key):
    return {r[key]: r for r in rows}


def test_clapi_parsing_and_inheritance():
    inv = parse_clapi_export(SAMPLE)
    assert len(inv.hosts) == 9
    web = inv.hosts["srv-lnx-web01"]
    assert web.poller == "Poller-DC1"
    assert web.templates == ["OS-Linux-SNMP-custom", "App-Webserver-Apache-ServerStatus-custom"]
    assert "generic-host" in inv.host_all_templates(web)
    assert not inv.hosts["sw-acc-02"].activate
    # services de groupes d'hôtes développés sur chaque membre
    pings = [s for s in inv.services if s.description == "Ping"]
    assert {s.host for s in pings} == {"srv-lnx-web01", "srv-lnx-web02", "srv-lnx-db01", "srv-lnx-erp01", "ups-dc1"}
    assert inv.stpl_chain("OS-Linux-Cpu-SNMP-custom") == [
        "OS-Linux-Cpu-SNMP-custom", "OS-Linux-Cpu-SNMP", "generic-active-service-custom",
        "generic-active-service", "generic-service"]
    assert inv.pack_template("OS-Linux-Cpu-SNMP-custom") == "OS-Linux-Cpu-SNMP"


def test_effective_check_resolves_macros_and_args():
    inv = parse_clapi_export(SAMPLE)
    svc = next(s for s in inv.services if s.host == "srv-lnx-db01" and s.description == "Cpu")
    chk = inv.effective_check(svc, inv.hosts["srv-lnx-db01"])
    assert chk.command == "OS-Linux-SNMP-Cpu"
    assert (chk.plugin, chk.mode) == ("os::linux::snmp::plugin", "cpu")
    assert "--warning-average='80'" in chk.line and "--snmp-community='public'" in chk.line
    batch = next(s for s in inv.services if s.description == "Batch-Mensuel")
    assert "-j MONTHLY -w 7200" in inv.effective_check(batch, inv.hosts["srv-lnx-erp01"]).line
    # macro mot de passe masquée à la lecture
    assert inv.host_templates["App-DB-MySQL-custom"].macros["MYSQLPASSWORD"] == "********"


def test_matrix_proposals():
    m = build_matrix(parse_clapi_export(SAMPLE), Mapper())
    stpl = _rows(m.service_templates, "Template de service")
    assert stpl["OS-Linux-Cpu-SNMP-custom"]["Proposition"] == MIGRER_NATIF
    assert stpl["OS-Linux-Cpu-SNMP-custom"]["Template Zabbix suggéré"] == "Linux by SNMP"
    assert stpl["OS-Linux-Ntp-SNMP-custom"]["Proposition"] == NE_PAS_MIGRER  # inutilisé
    assert stpl["generic-active-service-custom"]["Proposition"] == HERITE
    assert stpl["ERP-Batch-Nocturne"]["Stratégie"] == "Développer"
    assert stpl["Passive-Trap-Dummy"]["Proposition"] == NE_PAS_MIGRER
    assert "(aucun template) check_erp_batch" in stpl
    htpl = _rows(m.host_templates, "Template d'hôte")
    assert htpl["Net-Cisco-Standard-SNMP-custom"]["Template Zabbix suggéré"] == "Cisco IOS by SNMP"
    assert htpl["OS-Linux-SNMP-custom"]["Nb services issus du template"] == 12
    hosts = _rows(m.hosts, "Hôte")
    assert hosts["sw-acc-02"]["Proposition"] == NE_PAS_MIGRER
    assert "SNMP" in hosts["srv-lnx-web01"]["Interface Zabbix suggérée"]
    # aucun secret en clair
    assert all("public" not in r["Ligne de commande résolue"] for r in m.services)
    assert all("Sup3rS3cret" not in r["Ligne de commande résolue"] for r in m.services)


def test_mask_secrets():
    line = "check --password='abc' --snmp-community public -C secret --token=xyz"
    masked = mask_secrets(line)
    for secret in ("abc", "public", "secret", "xyz"):
        assert secret not in masked


def test_excel_and_csv(tmp_path):
    m = build_matrix(parse_clapi_export(SAMPLE))
    out = write_excel(m, tmp_path / "m.xlsx")
    wb = load_workbook(out)
    assert "Matrice templates services" in wb.sheetnames
    ws = wb["Services"]
    header = [c.value for c in ws[1]]
    formula = ws.cell(row=2, column=header.index("Décision (template)") + 1).value
    assert formula.startswith("=IFERROR(")
    assert len(write_csv(m, tmp_path / "csv")) == 5


def test_clapi_variants():
    inv = parse_clapi_lines([
        "CMD;ADD;c1;check;/bin/check -a $ARG1$",
        "HTPL;ADD;T1;T1;;;;",
        "HOST;ADD;h1;h1;1.1.1.1;;;",
        "HOST;addtemplate;h1;T1",
        "INSTANCE;addhost;Poller-X;h1",
        "HG;ADD;G;G",
        "HG;addmember;G;h1",
        "SERVICE;ADD;h1;s1;",
        "SERVICE;setparam;h1;s1;check_command;c1",
        "SERVICE;setparam;h1;s1;check_command_arguments;!42",
    ])
    h = inv.hosts["h1"]
    assert h.templates == ["T1"] and h.poller == "Poller-X" and h.hostgroups == ["G"]
    assert inv.effective_check(inv.services[0], h).line == "/bin/check -a 42"


def test_database_source_with_fake_rows():
    tables = {
        "FROM command": [{"command_id": 1, "command_name": "c1", "command_line": "x --mode=cpu", "command_type": 2}],
        "FROM host WHERE": [
            {"host_id": 1, "host_name": "T1", "host_alias": "", "host_address": "", "host_register": "0",
             "host_activate": "1", "command_command_id": None},
            {"host_id": 2, "host_name": "h1", "host_alias": "a", "host_address": "1.1.1.1", "host_register": "1",
             "host_activate": "1", "command_command_id": None}],
        "FROM host_template_relation": [{"host_host_id": 2, "host_tpl_id": 1}],
        "FROM ns_host_relation": [{"host_host_id": 2, "name": "Central"}],
        "FROM hostgroup_relation": [{"hostgroup_hg_id": 7, "host_host_id": 2}],
        "FROM hostgroup": [{"hg_id": 7, "hg_name": "G"}],
        "FROM on_demand_macro_host": [{"host_host_id": 2, "host_macro_name": "$_HOSTPWD$",
                                       "host_macro_value": "s3cr3t", "is_password": 1}],
        "FROM service WHERE": [
            {"service_id": 10, "service_description": "ST", "service_alias": "", "service_register": "0",
             "service_activate": "1", "service_template_model_stm_id": None, "command_command_id": 1,
             "command_command_id_arg": ""},
            {"service_id": 11, "service_description": "Cpu", "service_alias": "", "service_register": "1",
             "service_activate": "1", "service_template_model_stm_id": 10, "command_command_id": None,
             "command_command_id_arg": ""},
            {"service_id": 12, "service_description": "HGsvc", "service_alias": "", "service_register": "1",
             "service_activate": "0", "service_template_model_stm_id": 10, "command_command_id": None,
             "command_command_id_arg": ""}],
        "FROM on_demand_macro_service": [],
        "FROM host_service_relation": [
            {"host_host_id": 1, "hostgroup_hg_id": None, "service_service_id": 10},
            {"host_host_id": 2, "hostgroup_hg_id": None, "service_service_id": 11},
            {"host_host_id": None, "hostgroup_hg_id": 7, "service_service_id": 12}],
    }

    def query(sql):
        # la clé la plus longue d'abord : "FROM hostgroup_relation" avant "FROM hostgroup"
        for key in sorted(tables, key=len, reverse=True):
            if key in sql:
                return tables[key]
        raise AssertionError(sql)

    inv = build_inventory(query)
    h = inv.hosts["h1"]
    assert h.templates == ["T1"] and h.poller == "Central" and h.hostgroups == ["G"]
    assert h.macros["$_HOSTPWD$"] == "********"
    assert inv.host_templates["T1"].service_templates == ["ST"]
    assert {(s.description, s.activate, s.hostgroup) for s in inv.services} == {("Cpu", True, ""), ("HGsvc", False, "G")}
    assert inv.effective_check(inv.services[0], h).mode == "cpu"
