import io
import time
from pathlib import Path

import pytest
from openpyxl import load_workbook

from centreon_inventory.compare import ABSENT, PRESENT_A_LIER, PRESENT_OK, SVC_COUVERT
from centreon_inventory.filters import ClientFilter, detect_prefixes
from centreon_inventory.matrix import MIGRER_DEV
from centreon_inventory.report import build_report
from centreon_inventory.sources.clapi_export import parse_clapi_export
from centreon_inventory.zabbix import ZabbixError, load_zabbix

from .fake_centreon import FakeCentreon
from .fake_zabbix import FakeZabbix

SAMPLE = Path(__file__).resolve().parent.parent / "examples" / "centreon_export_exemple.txt"


@pytest.fixture(scope="module")
def inv():
    return parse_clapi_export(SAMPLE)


@pytest.fixture(scope="module")
def zbx():
    return load_zabbix("https://zabbix.test", token="zbx-token", session=FakeZabbix(), log=lambda m: None)


# ----------------------------------------------------------------------- filtre client
def test_client_filter_on_host_names(inv):
    f = ClientFilter(name="Linux", hosts="srv-LNX")  # insensible à la casse
    scoped = f.apply(inv)
    assert set(scoped.hosts) == {"srv-lnx-web01", "srv-lnx-web02", "srv-lnx-db01", "srv-lnx-erp01"}
    assert {s.host for s in scoped.services} <= set(scoped.hosts)


def test_client_filter_services_exclude_regex_and_or(inv):
    disks = ClientFilter(hosts="srv-lnx", services="Disk").apply(inv)
    assert {s.description for s in disks.services} == {"Disk-/", "Disk-/var/lib/mysql"}
    assert set(disks.hosts) == {"srv-lnx-web01", "srv-lnx-web02", "srv-lnx-db01"}
    excl = ClientFilter(hosts="srv-lnx", exclude="web02, erp").apply(inv)
    assert set(excl.hosts) == {"srv-lnx-web01", "srv-lnx-db01"}
    rx = ClientFilter(hosts="re:^sw-(core|acc)-").apply(inv)
    assert set(rx.hosts) == {"sw-core-01", "sw-acc-02"}
    either = ClientFilter(hosts="ups-", services="Trap", mode="ou").apply(inv)
    assert {(s.host, s.description) for s in either.services} == {("ups-dc1", "Ping"), ("srv-win-file01", "Trap-Sauvegarde")}


def test_client_report_hides_templates_unused_in_scope(inv):
    m = build_report(inv, client=ClientFilter(name="Réseau", hosts="sw-"))
    names = {r["Template de service"] for r in m.service_templates}
    assert "Net-Cisco-Standard-Environment-SNMP-custom" in names
    assert "OS-Linux-Cpu-SNMP-custom" not in names
    assert {r["Hôte"] for r in m.hosts} == {"sw-core-01", "sw-acc-02"}


def test_detect_prefixes():
    names = ["Mdph-973-mp-4225-EP", "Mdph-973-Srv-AdminVm", "Mdph-974-Srv-X", "Mdph-973-Srv-IwBdd", "solo"]
    assert detect_prefixes(names)[0] == ("Mdph-973", 3)


# ----------------------------------------------------------------------- Zabbix
def test_zabbix_inventory(zbx):
    assert zbx.version == "7.0.5" and len(zbx.hosts) == 5 and len(zbx.templates) == 9
    tpl = {t["name"]: t for t in zbx.templates}
    assert ["CPU utilization", "system.cpu.util[snmp.cpu.util]"] in tpl["Linux by SNMP"]["items"]
    assert any(k.startswith("vfs.fs.pused") for _, k in tpl["Linux by SNMP"]["prototypes"])
    assert tpl["Linux by Zabbix agent"]["linked"] == [tpl["Linux CPU by Zabbix agent"]["templateid"]]
    h = next(x for x in zbx.hosts if x["host"] == "srv-lnx-web01")
    assert h["proxy"] == "proxy-dc1" and "Apache by HTTP" in h["templates"] and h["items"] == 182


def test_zabbix_auth_variants():
    old = FakeZabbix(version="6.0.20")
    z = load_zabbix("https://z", user="Admin", password="zabbix", session=old, log=lambda m: None)
    assert len(z.hosts) == 5 and "user.login" in old.calls
    with pytest.raises(ZabbixError):
        load_zabbix("https://z", token="mauvais", session=FakeZabbix(), log=lambda m: None)


def test_compare_with_zabbix(inv, zbx):
    m = build_report(inv, zabbix=zbx)
    hosts = {r["Hôte Centreon"]: r for r in m.zabbix_hosts}
    assert hosts["srv-lnx-web01"]["Statut Zabbix"] == PRESENT_OK
    assert hosts["srv-lnx-db01"]["Statut Zabbix"] == PRESENT_A_LIER
    assert "MySQL by Zabbix agent 2" in hosts["srv-lnx-db01"]["Templates à lier"]
    assert hosts["srv-win-ad01"]["Correspondance"] == "Nom"  # casse différente dans Zabbix
    assert hosts["sw-core-01"]["Correspondance"] == "Adresse IP"
    assert hosts["srv-lnx-web02"]["Statut Zabbix"] == ABSENT
    assert [r["Hôte Zabbix"] for r in m.zabbix_only] == ["srv-lnx-new01"]
    svc = {(r["Hôte"], r["Service"]): r["Statut Zabbix"] for r in m.services}
    assert svc["srv-lnx-web01", "Cpu"] == SVC_COUVERT
    tpl = {r["Template Zabbix suggéré"]: r["Disponible dans Zabbix"] for r in m.zabbix_templates}
    assert tpl["Linux by SNMP"] == "Oui" and tpl["APC UPS by SNMP"].startswith("Non")


def test_decisions_flow_into_services_and_excel(inv, zbx, tmp_path):
    from centreon_inventory.excel import write_excel
    decisions = {"service_templates": {"ERP-Batch-Nocturne": {"Décision finale": MIGRER_DEV}},
                 "services": {"srv-lnx-db01|Cpu": {"Décision (service)": "Ne pas migrer"}}}
    m = build_report(inv, decisions=decisions, zabbix=zbx, client=ClientFilter(name="ERP", hosts="lnx"))
    svc = {(r["Hôte"], r["Service"]): r for r in m.services}
    assert svc["srv-lnx-erp01", "Batch-Nocturne"]["Décision effective"] == MIGRER_DEV
    assert svc["srv-lnx-db01", "Cpu"]["Décision effective"] == "Ne pas migrer"
    wb = load_workbook(write_excel(m, tmp_path / "x.xlsx"))
    assert {"Comparatif Zabbix hôtes", "Uniquement dans Zabbix", "Templates Zabbix"} <= set(wb.sheetnames)
    assert "Périmètre client : ERP" in wb["Synthèse"]["A3"].value


# ----------------------------------------------------------------------- interface web
@pytest.fixture()
def client(tmp_path, inv):
    from centreon_inventory.web.app import create_app
    app = create_app(tmp_path / "data", session_factory=FakeCentreon(inv), zabbix_session=FakeZabbix())
    app.config["TESTING"] = True
    return app.test_client()


def _wait(client, job):
    for _ in range(100):
        j = client.get(f"/api/jobs/{job}").get_json()
        if j["status"] != "en cours":
            return j
        time.sleep(0.05)
    raise AssertionError("job bloqué")


def test_web_end_to_end(client, tmp_path):
    assert client.get("/").status_code == 200
    assert client.get("/api/summary").get_json() == {"empty": True}
    # extraction Centreon par API puis Zabbix
    r = client.post("/api/jobs/centreon", json={"url": "https://c/centreon", "user": "admin", "password": "S3cret-Centreon"})
    assert _wait(client, r.get_json()["job"])["status"] == "terminé"
    r = client.post("/api/jobs/zabbix", json={"url": "https://z", "token": "zbx-token"})
    assert _wait(client, r.get_json()["job"])["status"] == "terminé"
    state = client.get("/api/state").get_json()
    assert state["centreon"]["hosts"] == 9 and state["zabbix"]["hosts"] == 5
    stored = "".join(p.read_text(encoding="utf-8") for p in (tmp_path / "data").glob("*.json"))
    assert "S3cret-Centreon" not in stored and "zbx-token" not in stored  # secrets jamais enregistrés
    # client enregistré + tableau filtré et trié
    client.post("/api/clients", json={"name": "Linux", "hosts": "srv-lnx"})
    t = client.get("/api/table/hosts?client=Linux&sort=Hôte&dir=desc").get_json()
    assert t["total"] == 4 and t["rows"][0]["Hôte"] == "srv-lnx-web02"
    t = client.get("/api/table/services?client=Linux&f.Service=disk").get_json()
    assert t["total"] == 3
    t = client.get("/api/table/zabbix_hosts?client=Linux&f.Statut Zabbix==Absent de Zabbix").get_json()
    assert {r["Hôte Centreon"] for r in t["rows"]} == {"srv-lnx-web02", "srv-lnx-erp01"}
    # décision saisie dans l'interface -> reportée sur les services
    r = client.post("/api/decision", json={"table": "service_templates", "key": "ERP-Batch-Nocturne",
                                           "field": "Décision finale", "value": MIGRER_DEV})
    assert r.status_code == 200
    t = client.get("/api/table/services?client=Linux&f.Service=Batch").get_json()
    assert {r["Décision effective"] for r in t["rows"]} == {MIGRER_DEV}
    assert client.post("/api/decision", json={"table": "services", "key": "a|b", "field": "Hôte",
                                              "value": "x"}).status_code == 400
    # export Excel du périmètre
    r = client.get("/api/export.xlsx?client=Linux")
    assert r.status_code == 200 and "Linux" in r.headers["Content-Disposition"]
    wb = load_workbook(io.BytesIO(r.data))
    hdr = [c.value for c in wb["Matrice templates services"][1]]
    rows = {row[0]: row for row in wb["Matrice templates services"].iter_rows(min_row=2, values_only=True)}
    assert rows["ERP-Batch-Nocturne"][hdr.index("Décision finale")] == MIGRER_DEV
    assert wb["Hôtes"].max_row == 5


def test_web_demo_and_clapi_import(client):
    assert client.post("/api/import/demo").status_code == 200
    s = client.get("/api/summary").get_json()
    assert s["kpi"]["hosts"] == 9 and s["kpi"]["zabbix"]
    data = {"file": (io.BytesIO(SAMPLE.read_bytes()), "export.txt")}
    r = client.post("/api/import/clapi", data=data, content_type="multipart/form-data")
    assert r.get_json()["hosts"] == 9


def test_web_basic_auth(tmp_path, monkeypatch):
    from centreon_inventory.web.app import create_app
    monkeypatch.setenv("WEB_USER", "u")
    monkeypatch.setenv("WEB_PASSWORD", "p")
    c = create_app(tmp_path / "d").test_client()
    assert c.get("/api/state").status_code == 401
    assert c.get("/api/state", headers={"Authorization": "Basic dTpw"}).status_code == 200


# ----------------------------------------------------------------------- retours de revue
def test_or_mode_keeps_matching_hosts_without_services(inv):
    from centreon_inventory.model import Host, Inventory
    inv2 = Inventory(commands=inv.commands, host_templates=inv.host_templates, service_templates=inv.service_templates,
                     hosts={**inv.hosts, "ups-sans-service": Host(name="ups-sans-service")}, services=list(inv.services))
    scoped = ClientFilter(hosts="ups-", services="Trap", mode="ou").apply(inv2)
    assert "ups-sans-service" in scoped.hosts and "srv-win-file01" in scoped.hosts


def test_zabbix_alias_and_ip_matches_are_unique(inv, zbx):
    import copy
    from centreon_inventory.model import Host
    inv2 = copy.deepcopy(inv)
    # deux hôtes Centreon avec le même alias, égal au nom d'un hôte Zabbix : pas de rapprochement ambigu
    inv2.hosts["dup-a"] = Host(name="dup-a", alias="srv-lnx-new01", address="10.9.9.1")
    inv2.hosts["dup-b"] = Host(name="dup-b", alias="srv-lnx-new01", address="10.9.9.2")
    # un hôte dont l'IP est celle d'un hôte Zabbix déjà rapproché par son nom
    inv2.hosts["copie-web01"] = Host(name="copie-web01", address="10.0.10.11")
    m = build_report(inv2, zabbix=zbx)
    hosts = {r["Hôte Centreon"]: r for r in m.zabbix_hosts}
    assert hosts["dup-a"]["Statut Zabbix"] == ABSENT and hosts["dup-b"]["Statut Zabbix"] == ABSENT
    assert hosts["copie-web01"]["Statut Zabbix"] == ABSENT
    assert hosts["srv-lnx-web01"]["Correspondance"] == "Nom"
    assert "srv-lnx-new01" in [r["Hôte Zabbix"] for r in m.zabbix_only]


def test_untrusted_values_are_never_formulas(inv, tmp_path):
    import csv
    from centreon_inventory.excel import write_csv, write_excel
    decisions = {"service_templates": {"Base-Ping-LAN": {"Commentaire": '=HYPERLINK("http://x","clic")'}}}
    m = build_report(inv, decisions=decisions)
    ws = load_workbook(write_excel(m, tmp_path / "x.xlsx"))["Matrice templates services"]
    hdr = [c.value for c in ws[1]]
    row = next(r for r in ws.iter_rows(min_row=2) if r[0].value == "Base-Ping-LAN")
    cell = row[hdr.index("Commentaire")]
    assert cell.data_type == "s" and cell.value.startswith("=HYPERLINK")
    write_csv(m, tmp_path / "csv")
    with open(tmp_path / "csv" / "matrice_templates_services.csv", encoding="utf-8-sig", newline="") as fh:
        rows = {r["Template de service"]: r for r in csv.DictReader(fh, delimiter=";")}
    assert rows["Base-Ping-LAN"]["Commentaire"].startswith("'=")


def test_web_scoped_centreon_collection(client):
    r = client.post("/api/jobs/centreon", json={"url": "https://c/centreon", "user": "a", "password": "p",
                                                "scope_only": True, "scope": {"hosts": ""}})
    assert r.status_code == 400
    r = client.post("/api/jobs/centreon", json={"url": "https://c/centreon", "user": "a", "password": "p",
                                                "scope_only": True, "scope": {"name": "Win", "hosts": "srv-win"}})
    j = _wait(client, r.get_json()["job"])
    assert j["status"] == "terminé", j
    state = client.get("/api/state").get_json()
    assert state["centreon"]["hosts"] == 2 and "collecte limitée" in state["centreon"]["source"]


def test_glob_hostgroup_and_poller_filters(inv):
    f = ClientFilter(hostgroups="HG-L*")
    assert set(f.apply(inv).hosts) == {"srv-lnx-web01", "srv-lnx-web02", "srv-lnx-db01", "srv-lnx-erp01"}
    f = ClientFilter(hostgroups="hg-web", pollers="Poller-DC1")
    assert set(f.apply(inv).hosts) == {"srv-lnx-web01", "srv-lnx-web02"}
    assert ClientFilter.from_dict(f.to_dict()) == f


def test_web_lookup_and_client_with_hostgroups(client):
    r = client.post("/api/centreon/lookup", json={"url": "https://c/centreon", "user": "a", "password": "p"})
    data = r.get_json()
    assert "HG-Linux" in data["hostgroups"] and "Poller-DC2" in data["pollers"]
    client.post("/api/clients", json={"name": "Web DC1", "hostgroups": "hg-web*", "pollers": "poller-dc1"})
    r = client.post("/api/jobs/centreon", json={"url": "https://c/centreon", "user": "a", "password": "p",
                                                "scope_only": True,
                                                "scope": {"name": "Web DC1", "hostgroups": "hg-web*", "pollers": "poller-dc1"}})
    assert _wait(client, r.get_json()["job"])["status"] == "terminé"
    t = client.get("/api/table/hosts?client=Web DC1").get_json()
    assert {row["Hôte"] for row in t["rows"]} == {"srv-lnx-web01", "srv-lnx-web02"}
    state = client.get("/api/state").get_json()
    assert dict(state["pollers"]) == {"Poller-DC1": 2}


def test_host_not_migrated_excludes_its_services(inv, tmp_path):
    from centreon_inventory.excel import write_excel
    decisions = {"hosts": {"srv-lnx-db01": {"Décision finale": "Ne pas migrer"}},
                 "services": {"srv-lnx-db01|Cpu": {"Décision (service)": MIGRER_DEV}}}
    m = build_report(inv, decisions=decisions)
    db = [r for r in m.services if r["Hôte"] == "srv-lnx-db01"]
    assert db and all(r["Décision effective"] == "Ne pas migrer" for r in db)
    assert all(r["Décision (hôte)"] == "Ne pas migrer" for r in db)
    other = next(r for r in m.services if r["Hôte"] == "srv-lnx-web01" and r["Service"] == "Cpu")
    assert other["Décision effective"] != "Ne pas migrer"
    # La formule Excel applique la même règle
    ws = load_workbook(write_excel(m, tmp_path / "h.xlsx"))["Services"]
    hdr = [c.value for c in ws[1]]
    f = ws.cell(row=2, column=hdr.index("Décision effective") + 1).value
    assert f.startswith('=IF(') and "Ne pas migrer" in f and "Hôtes" in ws.cell(row=2, column=hdr.index("Décision (hôte)") + 1).value


def test_web_bulk_decisions(client):
    client.post("/api/import/demo")
    keys = client.get("/api/table/service_templates/keys?f.Catégorie=CPU").get_json()["keys"]
    assert set(keys) == {"OS-Linux-Cpu-SNMP-custom", "OS-Windows-Cpu-SNMP-custom", "OS-Linux-Cpu-SNMP"}
    r = client.post("/api/decisions/bulk", json={"table": "service_templates", "keys": keys,
                                                 "field": "Décision finale", "value": "À analyser"})
    assert r.get_json()["count"] == 3
    t = client.get("/api/table/service_templates?f.Catégorie=CPU").get_json()
    assert {row["Décision finale"] for row in t["rows"]} == {"À analyser"}
    # hôtes : décision groupée « Ne pas migrer » -> leurs services ne sont plus migrés
    hk = client.get("/api/table/hosts/keys?f.Hôte=srv-win").get_json()["keys"]
    client.post("/api/decisions/bulk", json={"table": "hosts", "keys": hk, "field": "Décision finale",
                                             "value": "Ne pas migrer"})
    s = client.get("/api/table/services?f.Hôte=srv-win&size=100").get_json()
    assert s["rows"] and {row["Décision effective"] for row in s["rows"]} == {"Ne pas migrer"}
    # effacement groupé
    client.post("/api/decisions/bulk", json={"table": "hosts", "keys": hk, "field": "Décision finale", "value": ""})
    s = client.get("/api/table/services?f.Hôte=srv-win-ad01").get_json()
    assert "Ne pas migrer" not in {row["Décision effective"] for row in s["rows"] if row["Actif"] == "Oui"}
    assert client.post("/api/decisions/bulk", json={"table": "hosts", "keys": [], "field": "Décision finale"}).status_code == 400


def test_correlation_with_zabbix_items(inv, zbx):
    from centreon_inventory.correlate import A_IMPORTER, COUVERT, SANS_DONNEES, SPECIFIQUE
    m = build_report(inv, zabbix=zbx)
    corr = {r["Template de service"]: r for r in m.correlation}
    assert corr["OS-Linux-Cpu-SNMP-custom"]["Couverture standard"] == COUVERT
    assert "system.cpu.util" in corr["OS-Linux-Cpu-SNMP-custom"]["Items Zabbix correspondants"]
    assert "vfs.fs.pused" in corr["OS-Linux-Disk-Generic-Name-SNMP-custom"]["Items Zabbix correspondants"]  # LLD
    # suggestion « template constructeur » : on s'appuie sur le template Zabbix de l'hôte (Cisco)
    assert corr["Net-Cisco-Standard-Environment-SNMP-custom"]["Couverture standard"] == COUVERT
    # template suggéré absent de Zabbix -> à importer (pas de faux rapprochement avec Apache)
    assert corr["App-Protocol-HTTP-Response-Time-custom"]["Couverture standard"] == A_IMPORTER
    assert corr["ERP-Batch-Nocturne"]["Couverture standard"] == SPECIFIQUE
    assert corr["ERP-Batch-Nocturne"]["Approche recommandée"].startswith("Spécificité")
    svc = {(r["Hôte"], r["Service"]): r for r in m.services}
    assert svc["srv-lnx-web01", "Cpu"]["Couverture standard"] == COUVERT
    assert "CPU utilization" in svc["srv-lnx-web01", "Cpu"]["Item Zabbix standard"]
    tpl = {r["Template Zabbix suggéré"]: r for r in m.zabbix_templates}
    assert tpl["Linux by SNMP"]["Nb prototypes LLD"] == 2 and tpl["Linux by SNMP"]["Apport du standard (sans équivalent Centreon)"] >= 1
    # inventaire Zabbix enregistré par une version précédente (sans items)
    from centreon_inventory.zabbix import ZabbixInventory
    old = ZabbixInventory.from_dict({**zbx.to_dict(), "templates": [
        {k: v for k, v in t.items() if k not in ("items", "prototypes", "linked")} for t in zbx.templates]})
    m2 = build_report(inv, zabbix=old)
    assert {r["Couverture standard"] for r in m2.correlation} & {SANS_DONNEES}


def test_correlation_linked_templates(inv, zbx):
    """Les items des templates liés (ex. module CPU d'un template Linux) sont pris en compte."""
    from centreon_inventory.correlate import _Catalog
    cat = _Catalog(zbx)
    keys = {k for _, _, k, _ in cat.items(cat.get("Linux by Zabbix agent"))}
    assert "system.cpu.util" in keys and "vm.memory.util" in keys


def test_web_detail(client):
    client.post("/api/import/demo")
    d = client.get("/api/detail/service_templates?key=OS-Linux-Disk-Generic-Name-SNMP-custom").get_json()
    assert d["kind"] == "Template de service" and d["services"]["total"] == 3
    assert [c["name"] for c in d["chain"]][0] == "OS-Linux-Disk-Generic-Name-SNMP-custom"
    assert "--storage=" in d["example"]["line"]
    items = d["zabbix"][0]["items"]
    assert items[0]["match"] and items[0]["key"].startswith("vfs.fs")  # item correspondant en tête
    d = client.get("/api/detail/service_templates?key=(aucun template) check_erp_batch").get_json()
    assert d["no_template"] and d["services"]["rows"][0]["Hôte"] == "srv-lnx-erp01"
    assert client.get("/api/detail/host_templates?key=OS-Linux-SNMP-custom").get_json()["linked_service_templates"]
    assert client.get("/api/detail/hosts?key=srv-lnx-db01").get_json()["services"]["total"] == 6
    assert client.get("/api/detail/services?key=srv-lnx-db01|Cpu").status_code == 200
    assert client.get("/api/detail/services?key=inconnu|x").status_code == 404


def test_web_zabbix_template_autocomplete(client):
    # sans inventaire Zabbix : catalogue de la version la plus récente
    d = client.get("/api/zabbix/templates?context=IDRAC").get_json()
    assert d["items"][0]["name"] == "Dell iDRAC by SNMP" and d["catalog_version"] == d["versions"][-1]
    # avec l'inventaire Zabbix (7.0.5) : catalogue 7.0 + templates présents dans Zabbix
    client.post("/api/import/demo")
    d = client.get("/api/zabbix/templates?context=OS-Linux-SNMP").get_json()
    assert d["version"] == "7.0.5" and d["catalog_version"] == "7.0"
    assert d["items"][0]["name"] == "Linux by SNMP"
    # version visée choisie explicitement
    client.post("/api/settings", json={"zabbix_target_version": "6.0"})
    d = client.get("/api/zabbix/templates?q=interfaces").get_json()
    assert d["catalog_version"] == "6.0" and "Interfaces by SNMP" in [t["name"] for t in d["items"]]
    assert client.get("/api/state").get_json()["catalog"]["versions"]


def test_target_templates_cross_hosts_and_services(inv, zbx):
    decisions = {
        "hosts": {"srv-lnx-web01": {"Templates Zabbix cibles": "Linux by Zabbix agent, ICMP Ping"}},
        "host_templates": {"App-Webserver-Apache-ServerStatus-custom": {"Template Zabbix cible": "Apache by HTTP"}},
        "service_templates": {"OS-Linux-Cpu-SNMP-custom": {"Template / item Zabbix cible": "system.cpu.util[,idle]"}},
        "services": {"srv-lnx-web01|Cpu": {"Template Zabbix cible": "Linux by SNMP"}},
    }
    m = build_report(inv, decisions=decisions, zabbix=zbx)
    host = next(r for r in m.hosts if r["Hôte"] == "srv-lnx-web01")
    # choix de l'hôte + cible du template d'hôte + cible d'un service, sans doublon ni clé d'item
    assert host["Templates Zabbix retenus"].split(", ")[:3] == ["Linux by Zabbix agent", "ICMP Ping", "Apache by HTTP"]
    assert "Linux by SNMP" in host["Templates Zabbix retenus"]
    assert "system.cpu.util" not in host["Templates Zabbix retenus"]
    svc = {(r["Hôte"], r["Service"]): r for r in m.services}
    assert svc["srv-lnx-web01", "Cpu"]["Template Zabbix retenu"] == "Linux by SNMP"
    assert svc["srv-lnx-db01", "Cpu"]["Template Zabbix retenu"] == ""  # clé d'item : pas un template
    assert svc["srv-lnx-web01", "Ping"]["Templates Zabbix (hôte)"] == host["Templates Zabbix retenus"]
    zrow = next(r for r in m.zabbix_hosts if r["Hôte Centreon"] == "srv-lnx-web01")
    assert zrow["Templates Zabbix retenus"] == host["Templates Zabbix retenus"]
    assert zrow["Templates à lier"] == "Linux by Zabbix agent"  # le seul template retenu non lié dans Zabbix
    assert any(r["Template Zabbix suggéré"] == "Linux by Zabbix agent" for r in m.zabbix_templates)


def test_web_target_templates_editable(client):
    client.post("/api/import/demo", json={})
    r = client.post("/api/decision", json={"table": "hosts", "key": "srv-lnx-web01",
                                           "field": "Templates Zabbix cibles", "value": "ICMP Ping, Linux by SNMP"})
    assert r.status_code == 200
    t = client.get("/api/table/hosts?f.Hôte=srv-lnx-web01").get_json()
    assert "Templates Zabbix cibles" in t["editable"]
    assert t["rows"][0]["Templates Zabbix retenus"].startswith("ICMP Ping, Linux by SNMP")
    r = client.post("/api/decision", json={"table": "services", "key": "srv-lnx-web01|Cpu",
                                           "field": "Template Zabbix cible", "value": "Linux by Zabbix agent"})
    assert r.status_code == 200
    items = client.get("/api/zabbix/templates?suggested=Foo, ICMP Ping&q=icmp").get_json()["items"]
    assert items and items[0]["name"] == "ICMP Ping"


def test_host_excluded_templates_removed_from_cumul(inv):
    decisions = {
        "host_templates": {"App-Webserver-Apache-ServerStatus-custom": {"Template Zabbix cible": "Apache by HTTP"}},
        "hosts": {"srv-lnx-web01": {"Templates Zabbix cibles": "ICMP Ping",
                                    "Templates Zabbix exclus": "apache by http"}},
    }
    m = build_report(inv, decisions=decisions)
    host = next(r for r in m.hosts if r["Hôte"] == "srv-lnx-web01")
    assert host["Templates Zabbix retenus"] == "ICMP Ping"  # exclusion insensible à la casse
    svc = next(r for r in m.services if r["Hôte"] == "srv-lnx-web01")
    assert svc["Templates Zabbix (hôte)"] == "ICMP Ping"
