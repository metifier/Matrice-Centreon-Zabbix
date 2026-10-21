"""Écriture de la matrice au format Excel (.xlsx) et CSV."""
from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.worksheet.datavalidation import DataValidation

from .matrix import (A_ANALYSER, DECISIONS, HERITE, HOST_DECISIONS, MIGRER_ADAPTER, MIGRER_DEV,
                     MIGRER_NATIF, NE_PAS_MIGRER, SUMMARY_HEADERS, Matrix)
from .selection import DESACTIVER, ELEMENT_DECISIONS, FACULTATIF, SURVEILLER

SHEET_SUMMARY = "Synthèse"
SHEET_STPL = "Matrice templates services"
SHEET_HTPL = "Matrice templates hôtes"
SHEET_HOSTS = "Hôtes"
SHEET_SERVICES = "Services"
SHEET_X_HTPL = "Croisé hôtes x tpl hôtes"
SHEET_X_STPL = "Croisé hôtes x tpl services"
SHEET_CMD = "Commandes"
SHEET_HELP = "Mode d'emploi"
SHEET_ZBX_HOSTS = "Comparatif Zabbix hôtes"
SHEET_ZBX_ONLY = "Uniquement dans Zabbix"
SHEET_ZBX_TPL = "Templates Zabbix"
SHEET_CORR = "Corrélation items Zabbix"
SHEET_SEL = "Sélection items & problèmes"

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")  # colonnes à remplir par l'utilisateur
THIN = Side(style="thin", color="BFBFBF")
DECISION_COLORS = {
    MIGRER_NATIF: "C6EFCE", MIGRER_ADAPTER: "DDEBF7", MIGRER_DEV: "FCE4D6",
    NE_PAS_MIGRER: "D9D9D9", A_ANALYSER: "FFEB9C", HERITE: "EDEDED", "Migrer": "C6EFCE",
    SURVEILLER: "C6EFCE", FACULTATIF: "DDEBF7", DESACTIVER: "D9D9D9",
}
STATUS_COLORS = {
    "Présent – templates OK": "C6EFCE", "Couvert dans Zabbix": "C6EFCE", "Oui": None,
    "Présent – templates à lier": "FFEB9C", "Template à lier": "FFEB9C", "Présent": "DDEBF7",
    "Absent de Zabbix": "F8CBAD", "Hôte absent de Zabbix": "F8CBAD", "À créer dans Zabbix": "FCE4D6",
    "Non migré (décision)": "D9D9D9", "Non (à importer)": "F8CBAD",
}
STATUS_COLORS.update({"Couvert par un item standard": "C6EFCE", "Template standard sans item correspondant": "FFEB9C",
                      "Template standard à importer dans Zabbix": "DDEBF7", "Spécificité à créer": "FCE4D6",
                      "Non migré": "D9D9D9"})
STATUS_COLUMNS = {"Statut Zabbix", "Disponible dans Zabbix", "Couverture standard"}
INPUT_COLUMNS = {"Décision finale", "Décision (service)", "Template / item Zabbix cible",
                 "Template Zabbix cible", "Templates Zabbix cibles",
                 "Templates Zabbix exclus", "Commentaire", "Décision"}
MAX_WIDTH = {"Ligne de commande résolue": 80, "Ligne de commande": 90, "Chaîne d'héritage": 50,
             "Chaîne complète": 50, "Templates (tous)": 50}


def write_excel(m: Matrix, path) -> Path:
    """Écrit le classeur dans `path` (chemin ou flux binaire, ex. io.BytesIO)."""
    wb = Workbook()
    wb.remove(wb.active)
    _summary_sheet(wb, m)
    _table(wb, SHEET_STPL, m.service_templates, decisions=DECISIONS, key_width=40)
    _table(wb, SHEET_HTPL, m.host_templates, decisions=DECISIONS, key_width=35)
    _table(wb, SHEET_HOSTS, m.hosts, decisions=HOST_DECISIONS)
    ws = _table(wb, SHEET_SERVICES, m.services, decisions=DECISIONS)
    _service_formulas(ws, m)
    _cross(wb, SHEET_X_HTPL, *m.host_x_htpl, legend="D = template appliqué directement, h = hérité")
    _cross(wb, SHEET_X_STPL, *m.host_x_stpl, legend="Nombre de services de l'hôte utilisant ce template")
    if m.zabbix_hosts or m.zabbix_templates or m.zabbix_only:
        _table(wb, SHEET_ZBX_HOSTS, m.zabbix_hosts)
        _table(wb, SHEET_ZBX_ONLY, m.zabbix_only)
        _table(wb, SHEET_ZBX_TPL, m.zabbix_templates)
        if m.correlation:
            _table(wb, SHEET_CORR, m.correlation, key_width=40)
        if m.zabbix_elements:
            ws = _table(wb, SHEET_SEL, m.zabbix_elements, decisions=ELEMENT_DECISIONS, key_width=40)
            _element_formulas(ws, m)
    _table(wb, SHEET_CMD, m.commands)
    _help_sheet(wb)
    wb.calculation.fullCalcOnLoad = True
    if hasattr(path, "write"):
        wb.save(path)
        return path
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def write_csv(m: Matrix, directory: str | Path) -> list[Path]:
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    written = []
    for name, rows in (("matrice_templates_services", m.service_templates),
                       ("matrice_templates_hotes", m.host_templates), ("hotes", m.hosts),
                       ("services", m.services), ("commandes", m.commands),
                       ("comparatif_zabbix_hotes", m.zabbix_hosts), ("uniquement_zabbix", m.zabbix_only),
                       ("templates_zabbix", m.zabbix_templates), ("correlation_items_zabbix", m.correlation),
                       ("selection_items_problemes_zabbix", m.zabbix_elements)):
        if not rows and name.startswith(("comparatif", "uniquement", "templates_zabbix", "correlation", "selection")):
            continue
        p = directory / f"{name}.csv"
        with p.open("w", newline="", encoding="utf-8-sig") as fh:
            if rows:
                w = csv.DictWriter(fh, fieldnames=list(rows[0]), delimiter=";")
                w.writeheader()
                w.writerows({k: _csv_safe(v) for k, v in r.items()} for r in rows)
        written.append(p)
    return written


def _csv_safe(v):
    """Neutralise les valeurs qu'Excel interpréterait comme une formule à l'ouverture du CSV."""
    if v is None:
        return ""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v


# --------------------------------------------------------------------------- helpers
def _table(wb, title, rows, decisions=None, key_width=None):
    ws = wb.create_sheet(title)
    if not rows:
        ws.append(["(aucune donnée)"])
        return ws
    headers = list(rows[0])
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h) for h in headers])
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"  # donnée Centreon / saisie utilisateur : jamais une formule
    _style_header(ws, len(headers))
    n = len(rows) + 1
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{n}"
    for idx, h in enumerate(headers, 1):
        letter = get_column_letter(idx)
        values = [len(str(r.get(h) or "")) for r in rows[:500]]
        width = min(max([len(h) + 2] + values) + 2, MAX_WIDTH.get(h, 45))
        if idx == 1 and key_width:
            width = max(width, key_width)
        ws.column_dimensions[letter].width = width
        if h in INPUT_COLUMNS:
            for row in ws.iter_rows(min_row=2, max_row=n, min_col=idx, max_col=idx):
                row[0].fill = INPUT_FILL
        if decisions and h in ("Décision finale", "Décision (service)", "Décision"):
            dv = DataValidation(type="list", formula1='"' + ",".join(decisions) + '"', allow_blank=True,
                                showErrorMessage=True, errorTitle="Décision",
                                error="Choisir une valeur dans la liste")
            ws.add_data_validation(dv)
            dv.add(f"{letter}2:{letter}{n}")
        if h in STATUS_COLUMNS:
            for status, color in STATUS_COLORS.items():
                if color:
                    ws.conditional_formatting.add(
                        f"{letter}2:{letter}{n}",
                        FormulaRule(formula=[f'{letter}2="{status}"'], fill=PatternFill("solid", fgColor=color)))
        if h in ("Proposition", "Décision finale", "Décision (template)", "Décision (hôte)", "Décision effective",
                 "Décision (service)", "Décision"):
            for decision, color in DECISION_COLORS.items():
                ws.conditional_formatting.add(
                    f"{letter}2:{letter}{n}",
                    FormulaRule(formula=[f'{letter}2="{decision}"'], fill=PatternFill("solid", fgColor=color)))
    return ws


def _style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=THIN)
    ws.row_dimensions[1].height = 32


def _service_formulas(ws, m: Matrix):
    """Les services héritent de la décision prise sur leur template (modifiable ligne par ligne)."""
    if not m.services or not m.service_templates:
        return
    headers = list(m.services[0])
    stpl_headers = list(m.service_templates[0])
    col = {h: get_column_letter(i) for i, h in enumerate(headers, 1)}
    s_col = {h: get_column_letter(i) for i, h in enumerate(stpl_headers, 1)}
    sheet = quote_sheetname(SHEET_STPL)
    last = len(m.service_templates) + 1
    keys = f"{sheet}!${s_col['Template de service']}$2:${s_col['Template de service']}${last}"
    final = f"{sheet}!${s_col['Décision finale']}$2:${s_col['Décision finale']}${last}"
    proposal = f"{sheet}!${s_col['Proposition']}$2:${s_col['Proposition']}${last}"
    hosts_ref = None
    if m.hosts:
        h_col = {h: get_column_letter(i) for i, h in enumerate(m.hosts[0], 1)}
        hsheet, hlast = quote_sheetname(SHEET_HOSTS), len(m.hosts) + 1
        rng = lambda c: f"{hsheet}!${h_col[c]}$2:${h_col[c]}${hlast}"  # noqa: E731
        hosts_ref = (rng("Hôte"), rng("Décision finale"), rng("Proposition"))
    for r in range(2, len(m.services) + 2):
        key = f"${col['Template de service']}{r}"
        match = f"MATCH({key},{keys},0)"
        ws[f"{col['Décision (template)']}{r}"] = (
            f'=IFERROR(IF(INDEX({final},{match})<>"",INDEX({final},{match}),INDEX({proposal},{match})),"")')
        svc_dec, tpl_dec, actif = (f"{col['Décision (service)']}{r}", f"{col['Décision (template)']}{r}",
                                   f"{col['Actif']}{r}")
        host_dec = f"{col['Décision (hôte)']}{r}"
        if hosts_ref:
            hkeys, hfinal, hprop = hosts_ref
            hmatch = f"MATCH(${col['Hôte']}{r},{hkeys},0)"
            ws[host_dec] = (f'=IFERROR(IF(INDEX({hfinal},{hmatch})<>"",INDEX({hfinal},{hmatch}),'
                            f'INDEX({hprop},{hmatch})),"")')
        ws[f"{col['Décision effective']}{r}"] = (
            f'=IF({host_dec}="{NE_PAS_MIGRER}","{NE_PAS_MIGRER}",'
            f'IF({svc_dec}<>"",{svc_dec},IF({actif}="Non","{NE_PAS_MIGRER}",{tpl_dec})))')


def _element_formulas(ws, m: Matrix):
    """Items / problèmes Zabbix : la décision effective suit la « Décision » saisie, sinon la proposition.
    Un problème désactivé parce que tous ses items le sont (colonne « Cohérence ») garde sa valeur calculée."""
    col = {h: get_column_letter(i) for i, h in enumerate(m.zabbix_elements[0], 1)}
    for r, row in enumerate(m.zabbix_elements, 2):
        if row["Cohérence"].startswith("Tous ses items"):
            continue
        dec, prop = f"{col['Décision']}{r}", f"{col['Proposition']}{r}"
        ws[f"{col['Décision effective']}{r}"] = f'=IF({dec}<>"",{dec},{prop})'


def _cross(wb, title, columns, rows, legend):
    ws = wb.create_sheet(title)
    ws.append([legend])
    ws["A1"].font = Font(italic=True)
    ws.append(["Hôte"] + list(columns))
    for r in rows:
        ws.append([r["Hôte"]] + [r.get(c, "") for c in columns])
    for c in range(1, len(columns) + 2):
        cell = ws.cell(row=2, column=c)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        if c > 1:
            cell.alignment = Alignment(text_rotation=90, horizontal="center")
            ws.column_dimensions[get_column_letter(c)].width = 4
    ws.row_dimensions[2].height = 200
    ws.column_dimensions["A"].width = max([10] + [len(r["Hôte"]) + 2 for r in rows])
    ws.freeze_panes = "B3"
    if columns and rows:
        rng = f"B3:{get_column_letter(len(columns) + 1)}{len(rows) + 2}"
        ws.conditional_formatting.add(rng, FormulaRule(formula=['B3="D"'], fill=PatternFill("solid", fgColor="8EA9DB")))
        ws.conditional_formatting.add(rng, FormulaRule(formula=['AND(B3<>"",B3<>"D")'],
                                                       fill=PatternFill("solid", fgColor="D9E1F2")))
        for row in ws.iter_rows(min_row=3, min_col=2):
            for cell in row:
                cell.alignment = Alignment(horizontal="center")


def _summary_sheet(wb, m: Matrix):
    ws = wb.create_sheet(SHEET_SUMMARY)
    ws.append(["Inventaire Centreon – matrice de migration vers Zabbix"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Généré le {datetime.now():%d/%m/%Y %H:%M}"])
    client = getattr(m, "client", None)
    if client is not None and client.active:
        ws.append([f"Périmètre client : {client.name or '(filtre)'} — hôtes : {client.hosts or '*'} ; "
                   f"services : {client.services or '*'} ; exclusions : {client.exclude or 'aucune'}"])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color="C00000")
    ws.append([])
    for label, value in list(m.summary) + ([(None, None)] + list(m.zabbix_summary) if m.zabbix_summary else []):
        ws.append([label, value])
        if label in SUMMARY_HEADERS:
            for c in (1, 2):
                cell = ws.cell(row=ws.max_row, column=c)
                cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    ws.append([])
    ws.append(["Top 15 des templates de service (par nombre de services)"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.append(["Template de service", "Nb services", "Nb hôtes", "Proposition", "Template Zabbix suggéré"])
    for c in range(1, 6):
        cell = ws.cell(row=ws.max_row, column=c)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for r in m.service_templates[:15]:
        if r["Nb services"]:
            ws.append([r["Template de service"], r["Nb services"], r["Nb hôtes"], r["Proposition"],
                       r["Template Zabbix suggéré"]])
    ws.column_dimensions["A"].width = 50
    ws.column_dimensions["B"].width = 26
    for col in "CDE":
        ws.column_dimensions[col].width = 30


def _help_sheet(wb):
    ws = wb.create_sheet(SHEET_HELP)
    lines = [
        ("Mode d'emploi de la matrice", True),
        ("", False),
        ("1. « Matrice templates services » : une ligne par template de service Centreon "
         "(et par commande pour les services sans template). C'est l'unité de décision principale.", False),
        ("   Lire la colonne « Proposition » (calculée), puis saisir la « Décision finale » (liste déroulante) "
         "et le « Template / item Zabbix cible ». Les cellules jaunes sont à remplir.", False),
        ("2. « Matrice templates hôtes » : même principe pour les templates d'hôte → templates Zabbix.", False),
        ("3. « Services » : inventaire détaillé hôte par hôte. La colonne « Décision (template) » reprend "
         "automatiquement la décision du template ; « Décision (service) » permet une exception ponctuelle ; "
         "« Décision effective » donne le résultat final (filtrable).", False),
        ("4. « Hôtes » : liste des hôtes avec interface et templates Zabbix suggérés, proxy = collecteur Centreon. "
         "« Templates Zabbix cibles » : templates choisis pour l'hôte (plusieurs, séparés par des virgules).", False),
        ("   « Templates Zabbix retenus » croise ces choix avec les cibles des templates d'hôte et des services migrés "
         "de l'hôte, moins les « Templates Zabbix exclus » (cumul erroné à retirer) ; côté « Services », « Template Zabbix cible » (exception) l'emporte sur la cible du template de "
         "service, et « Templates Zabbix (hôte) » rappelle les templates retenus pour l'hôte.", False),
        ("5. Onglets « Croisé » : vue matricielle hôtes × templates pour repérer les profils d'hôtes.", False),
        ("6. « Corrélation items Zabbix » : pour chaque template de service Centreon, item(s) des templates Zabbix "
         "standards qui couvrent déjà la mesure, ou spécificité à créer. Démarche conseillée : lier d'abord les "
         "templates standards, puis n'ajouter que les spécificités.", False),
        ("7. « Sélection items & problèmes » : items, prototypes LLD et problèmes (triggers) des templates Zabbix "
         "retenus pour les hôtes migrés, plus les items à créer (spécificités). Proposition : « Surveiller » "
         "(équivalent d'un service Centreon, problème grave ou item utilisé par un problème surveillé), "
         "« Facultatif » (apport du standard), « Désactiver » (bruit : faible sévérité, changement d'inventaire). "
         "Saisir la « Décision » ; un problème dont tous les items sont désactivés est désactivé (« Cohérence »).",
         False),
        ("", False),
        ("Stratégies", True),
        ("Natif : couvert par un template officiel Zabbix → effort faible.", False),
        ("Adapter : équivalent Zabbix existant (clé d'item, check simple, LLD) à paramétrer → effort moyen.", False),
        ("Développer : check spécifique (script maison) → UserParameter / script externe / plugin agent 2.", False),
        ("Abandonner : sans objet sous Zabbix (dummy, supervision de Centreon lui-même...).", False),
        ("", False),
        ("Priorité : Haute / Moyenne / Basse selon le nombre d'hôtes concernés (seuils paramétrables).", False),
        ("Les correspondances proviennent de mapping/zabbix_mapping.yaml : adaptez-le puis régénérez la matrice.", False),
        ("Les secrets détectés dans les lignes de commande (--password, community SNMP...) sont masqués.", False),
    ]
    for text, bold in lines:
        ws.append([text])
        if bold:
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
    ws.column_dimensions["A"].width = 140
